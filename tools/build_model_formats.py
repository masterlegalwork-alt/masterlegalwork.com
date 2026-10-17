"""Build the 28 published Word models from their readable HTML source.

The MLW house layout is a model layout, not a certification of registry compliance.
Render every changed document and refresh PDF counterparts before deployment.
"""
from pathlib import Path
import re, json, shutil
from lxml import html
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]

def text(el):
    return ' '.join(el.text_content().split())

def add(d, value, heading=False, align=None, keep=False):
    p=d.add_paragraph(); p.alignment=align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.keep_with_next=keep or heading
    p.paragraph_format.widow_control=True
    r=p.add_run(value);r.bold=heading;r.underline=heading
    return p

def build(node, destination):
    d=Document();s=d.sections[0]
    s.page_width=Inches(8.5);s.page_height=Inches(14)
    s.left_margin=Inches(1.6);s.right_margin=Inches(1.2)
    s.top_margin=s.bottom_margin=Inches(1.2)
    for style in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
        st=d.styles[style];st.font.name='Arial';st.font.size=Pt(13);st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.line_spacing=1.25;st.paragraph_format.space_after=Pt(4)
    page=s.header.paragraphs[0];page.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');page._p.append(fld)
    d.core_properties.author='Master Legal Work';d.core_properties.subject='Model pleading in the MLW house layout'
    d.core_properties.comments='Layout revised 17 October 2026. Adapt to facts and mandatory current court rules.'
    affidavit=False
    def walk(el):
        nonlocal affidavit
        cls=el.get('class','');value=text(el)
        if cls=='doc-print-note':return
        if el.tag=='hr':
            affidavit=True
            return
        if el.tag=='section' and 'draft-guidance' in cls:
            p=add(d,'PREPARATION AND JUDGMENT APPLICATION',True,WD_ALIGN_PARAGRAPH.CENTER)
            p.paragraph_format.page_break_before=True
            add(d,'MODEL FOR GENERAL REFERENCE ONLY. NOT LEGAL ADVICE. These preparation notes are not part of the document to be signed or filed.')
            for child in el:
                if child.tag in ['h3','h4']:continue
                walk(child)
            return
        if el.tag=='ol':
            start=int(el.get('start','1'))
            for i,li in enumerate(el,start):
                p=add(d,f'{i}. {text(li)}');p.paragraph_format.left_indent=Inches(.25);p.paragraph_format.first_line_indent=Inches(-.25)
            return
        if cls=='d-sig':
            if destination.stem in ['anticipatory-bail','execution-order-21']:
                d.paragraphs[-1].paragraph_format.keep_with_next=True
            for child in el:
                walk(child)
                if destination.stem in ['anticipatory-bail','execution-order-21']:d.paragraphs[-1].paragraph_format.keep_with_next=True
            return
        if cls=='counsel-block':
            # Four deliberate lines, never a flattened run of names and labels.
            lines=[x.strip() for x in el.itertext() if x.strip()]
            d.paragraphs[-1].paragraph_format.keep_with_next=True
            p=add(d,'\n'.join(lines),align=WD_ALIGN_PARAGRAPH.RIGHT)
            p.paragraph_format.keep_together=True
            return
        if not value:return
        if cls=='d-party':
            role=el.xpath('./span[@class="role"]')
            content=el.xpath('./span[not(@class="role")]')
            add(d,text(content[0]) if content else value,keep=bool(role))
            if role:add(d,text(role[0]),align=WD_ALIGN_PARAGRAPH.RIGHT)
            return
        heading=cls in ['d-court','d-title','d-h'] or el.tag in ['h3','h4']
        align=WD_ALIGN_PARAGRAPH.CENTER if heading or cls in ['d-vs','d-caseno'] else WD_ALIGN_PARAGRAPH.RIGHT if cls=='d-right' else None
        p=add(d,value,heading,align,keep=cls in ['d-caseno','d-vs'])
        if affidavit:p.paragraph_format.space_after=Pt(2)
        if destination.stem=='execution-order-21' and cls=='d-h' and value=='VERIFICATION' and not affidavit:
            p.paragraph_format.page_break_before=True
        if cls=='d-court' and affidavit:
            p.paragraph_format.page_break_before=True
        if cls=='d-verification' or 'Verified at' in value:
            p.paragraph_format.keep_together=True
            p.paragraph_format.keep_with_next=True
        if cls=='d-right':p.paragraph_format.keep_with_next=True
        for link in el.xpath('.//a[@href]'):
            add(d,'Source: '+link.get('href'))
    for el in node:walk(el)
    destination.parent.mkdir(parents=True,exist_ok=True);d.save(destination)

def main():
    source=html.fromstring((ROOT/'drafts.html').read_text())
    out=[]
    for article in source.xpath('//article[@class="draft"]'):
        key=article.get('id');node=article.xpath('.//div[@class="doc"]')[0]
        target=ROOT/'assets/drafts'/f'{key}.docx';build(node,target);out.append(str(target))
    print(json.dumps(out))

if __name__=='__main__':main()
